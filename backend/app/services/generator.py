"""Generate DOCX and PDF output files from a ResumeIR."""

from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path
from typing import Optional

from app.models.ir import ResumeIR


# ---------------------------------------------------------------------------
# DOCX Generator
# ---------------------------------------------------------------------------

def generate_docx(ir: ResumeIR) -> bytes:
    """Convert a ResumeIR to a DOCX file and return raw bytes."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Remove default empty paragraph
    for para in doc.paragraphs:
        p = para._element
        p.getparent().remove(p)

    # --- Contact info ---
    contact = ir.contact_info
    if contact.get("name"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(contact["name"])
        run.bold = True
        run.font.size = Pt(16)

    contact_parts = [
        v for k, v in contact.items() if k != "name" and v
    ]
    if contact_parts:
        p = doc.add_paragraph(" | ".join(contact_parts))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- Sections ---
    for section in ir.sections:
        # Section heading
        p = doc.add_paragraph()
        run = p.add_run(section.heading.upper())
        run.bold = True
        font_size = section.heading_style.get("font_size")
        if font_size:
            run.font.size = Pt(float(font_size))
        else:
            run.font.size = Pt(12)

        # Horizontal rule via bottom border on the heading paragraph
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        pBdr.append(bottom)
        pPr.append(pBdr)

        if section.section_type in ("skills", "summary"):
            for item in section.plain_items:
                p = doc.add_paragraph(style="List Bullet")
                _add_runs_to_para(p, item)
        else:
            for entry in section.entries:
                # Entry header line
                header_parts = []
                if entry.employer:
                    header_parts.append(entry.employer)
                if entry.title:
                    header_parts.append(entry.title)

                date_parts = []
                if entry.start_date:
                    date_parts.append(entry.start_date)
                if entry.end_date:
                    date_parts.append(entry.end_date)

                if header_parts or date_parts:
                    p = doc.add_paragraph()
                    if header_parts:
                        run = p.add_run(" | ".join(header_parts))
                        run.bold = True
                    if date_parts:
                        if header_parts:
                            p.add_run("  ")
                        run = p.add_run(" – ".join(date_parts))
                        run.italic = True
                    if entry.location:
                        p.add_run(f"  {entry.location}")

                for bullet in entry.bullets:
                    p = doc.add_paragraph(style="List Bullet")
                    _add_runs_to_para(p, bullet)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_runs_to_para(para: object, item: object) -> None:
    """Add TextRun objects from a BulletItem to a python-docx paragraph."""
    from docx.shared import Pt

    if item.runs:  # type: ignore[union-attr]
        for tr in item.runs:  # type: ignore[union-attr]
            if not tr.text:
                continue
            run = para.add_run(tr.text)  # type: ignore[union-attr]
            run.bold = tr.bold
            run.italic = tr.italic
            run.underline = tr.underline
            if tr.font_size:
                run.font.size = Pt(tr.font_size)
    else:
        para.add_run(item.text)  # type: ignore[union-attr]


# ---------------------------------------------------------------------------
# PDF Generator
# ---------------------------------------------------------------------------

def generate_pdf(ir: ResumeIR) -> bytes:
    """Convert a ResumeIR to PDF bytes using reportlab."""
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        HRFlowable,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name",
        parent=styles["Normal"],
        fontSize=18,
        fontName="Helvetica-Bold",
        alignment=1,  # center
        spaceAfter=4,
    )
    contact_style = ParagraphStyle(
        "Contact",
        parent=styles["Normal"],
        fontSize=10,
        alignment=1,
        spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "Heading",
        parent=styles["Normal"],
        fontSize=12,
        fontName="Helvetica-Bold",
        spaceBefore=10,
        spaceAfter=4,
    )
    entry_header_style = ParagraphStyle(
        "EntryHeader",
        parent=styles["Normal"],
        fontSize=10,
        fontName="Helvetica-Bold",
        spaceBefore=4,
        spaceAfter=2,
    )
    bullet_style = ParagraphStyle(
        "Bullet",
        parent=styles["Normal"],
        fontSize=10,
        leftIndent=18,
        bulletIndent=6,
        spaceAfter=2,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=10,
        spaceAfter=2,
    )

    story = []

    contact = ir.contact_info
    if contact.get("name"):
        story.append(Paragraph(contact["name"], name_style))

    contact_parts = [v for k, v in contact.items() if k != "name" and v]
    if contact_parts:
        story.append(Paragraph(" | ".join(contact_parts), contact_style))

    for section in ir.sections:
        story.append(Paragraph(section.heading.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.black))

        if section.section_type in ("skills", "summary"):
            for item in section.plain_items:
                story.append(Paragraph(f"• {item.text}", bullet_style))
        else:
            for entry in section.entries:
                header_parts = []
                if entry.employer:
                    header_parts.append(f"<b>{entry.employer}</b>")
                if entry.title:
                    header_parts.append(entry.title)

                date_parts = []
                if entry.start_date:
                    date_parts.append(entry.start_date)
                if entry.end_date:
                    date_parts.append(entry.end_date)

                if header_parts or date_parts:
                    header_text = " | ".join(header_parts)
                    if date_parts:
                        header_text += f"  <i>{' – '.join(date_parts)}</i>"
                    if entry.location:
                        header_text += f"  {entry.location}"
                    story.append(Paragraph(header_text, entry_header_style))

                for bullet in entry.bullets:
                    story.append(Paragraph(f"• {bullet.text}", bullet_style))

        story.append(Spacer(1, 6))

    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Convenience
# ---------------------------------------------------------------------------

def generate_file(ir: ResumeIR, fmt: str) -> bytes:
    """Generate a file in the requested format ('docx' or 'pdf')."""
    if fmt == "docx":
        return generate_docx(ir)
    elif fmt == "pdf":
        return generate_pdf(ir)
    else:
        raise ValueError(f"Unsupported format: {fmt}")
