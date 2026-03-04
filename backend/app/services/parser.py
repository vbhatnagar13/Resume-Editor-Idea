"""Resume file parser: converts DOCX and PDF files into ResumeIR."""

from __future__ import annotations

import re
import uuid
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from app.models.ir import BulletItem, ResumeEntry, ResumeIR, ResumeSection, TextRun

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SECTION_KEYWORDS: Dict[str, str] = {
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "work history": "experience",
    "education": "education",
    "academic": "education",
    "skills": "skills",
    "technical skills": "skills",
    "core competencies": "skills",
    "competencies": "skills",
    "technologies": "skills",
    "summary": "summary",
    "objective": "summary",
    "profile": "summary",
    "professional summary": "summary",
    "career summary": "summary",
    "about": "summary",
}

BULLET_CHARS = {"•", "‣", "◦", "▪", "▸", "-", "–", "—", "*"}

DATE_PATTERN = re.compile(
    r"""
    (?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|
       Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)
    [\s,]*\d{4}
    |
    \d{1,2}/\d{4}
    |
    \d{4}\s*[-–]\s*(?:\d{4}|[Pp]resent|[Cc]urrent|[Nn]ow)
    """,
    re.VERBOSE | re.IGNORECASE,
)


# ---------------------------------------------------------------------------
# DOCX Parser
# ---------------------------------------------------------------------------

def _run_to_textrun(run: Any) -> TextRun:
    """Convert a python-docx Run into a TextRun."""
    font_size: Optional[float] = None
    if run.font.size is not None:
        font_size = run.font.size.pt  # type: ignore[attr-defined]
    return TextRun(
        text=run.text,
        bold=bool(run.bold),
        italic=bool(run.italic),
        underline=bool(run.underline),
        font_size=font_size,
        font_name=run.font.name,
    )


def _para_font_size(para: Any) -> Optional[float]:
    """Return the dominant font size of a paragraph."""
    sizes = [r.font.size.pt for r in para.runs if r.font.size]
    if sizes:
        return max(sizes)
    return None


def _is_heading(para: Any, base_size: float) -> bool:
    """Heuristic: paragraph is a section heading."""
    text = para.text.strip()
    if not text:
        return False
    # Named heading style
    if para.style and "Heading" in (para.style.name or ""):
        return True
    # ALL CAPS short line
    if text.isupper() and len(text) < 60:
        return True
    # Bold text + larger than base size, or bold with short text
    size = _para_font_size(para)
    all_bold = all(r.bold for r in para.runs if r.text.strip())
    if all_bold and para.runs:
        if size and base_size and size > base_size:
            return True
        if len(text) < 50:
            return True
    return False


def _detect_base_font_size(paragraphs: List[Any]) -> float:
    """Detect the most common font size in the document."""
    from collections import Counter
    sizes: List[float] = []
    for para in paragraphs:
        for run in para.runs:
            if run.font.size:
                sizes.append(run.font.size.pt)
    if not sizes:
        return 11.0
    counter = Counter(sizes)
    return counter.most_common(1)[0][0]


def _strip_bullet_char(text: str) -> str:
    """Remove leading bullet characters."""
    text = text.strip()
    if text and text[0] in BULLET_CHARS:
        text = text[1:].strip()
    return text


def _make_bullet(text: str, runs: List[TextRun], level: int = 0) -> BulletItem:
    clean = _strip_bullet_char(text)
    return BulletItem(
        id=str(uuid.uuid4()),
        text=clean,
        runs=runs,
        level=level,
        original_text=clean,
    )


def _classify_section(heading: str) -> str:
    key = heading.strip().lower().rstrip(":")
    return SECTION_KEYWORDS.get(key, "other")


def _extract_employer_title_dates(
    lines: List[str],
) -> Tuple[Optional[str], Optional[str], Optional[str], Optional[str], Optional[str]]:
    """
    Naive heuristic: first non-empty line is employer or title,
    look for date ranges with regex.
    """
    employer: Optional[str] = None
    title: Optional[str] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    location: Optional[str] = None

    non_empty = [l.strip() for l in lines if l.strip()]
    if not non_empty:
        return employer, title, start_date, end_date, location

    # Look for date spans
    all_text = " ".join(non_empty)
    date_matches = DATE_PATTERN.findall(all_text)

    if len(date_matches) >= 2:
        start_date, end_date = date_matches[0], date_matches[1]
    elif len(date_matches) == 1:
        start_date = date_matches[0]

    # First two non-date, non-empty lines → employer, title
    info_lines = [
        l for l in non_empty if not DATE_PATTERN.search(l)
    ]
    if info_lines:
        employer = info_lines[0]
    if len(info_lines) >= 2:
        title = info_lines[1]

    return employer, title, start_date, end_date, location


def parse_docx(file_bytes: bytes) -> ResumeIR:
    """Parse a DOCX file into a ResumeIR."""
    import docx  # python-docx

    doc = docx.Document(BytesIO(file_bytes))

    has_tables = len(doc.tables) > 0
    paragraphs = doc.paragraphs
    base_size = _detect_base_font_size(paragraphs)

    sections: List[ResumeSection] = []
    current_section: Optional[ResumeSection] = None
    current_entry_lines: List[str] = []
    current_entry_paras: List[Any] = []
    contact_lines: List[str] = []
    found_first_section = False

    def flush_entry() -> None:
        """Finalize the current entry and add to section."""
        nonlocal current_entry_lines, current_entry_paras
        if not current_section or not current_entry_lines:
            current_entry_lines = []
            current_entry_paras = []
            return

        bullet_paras = [
            p for p in current_entry_paras
            if p.text.strip() and (
                p.style.name.startswith("List") or
                (p.text.strip() and p.text.strip()[0] in BULLET_CHARS) or
                (p.paragraph_format.left_indent and p.paragraph_format.left_indent > 0)
            )
        ]
        header_lines = [
            l for l in current_entry_lines
            if l.strip() and not any(
                p.text.strip() == l.strip() and (
                    p.style.name.startswith("List") or
                    (p.text.strip() and p.text.strip()[0] in BULLET_CHARS)
                )
                for p in current_entry_paras
            )
        ]

        employer, title, start_date, end_date, location = _extract_employer_title_dates(
            header_lines
        )

        bullets: List[BulletItem] = []
        for bp in bullet_paras:
            runs = [_run_to_textrun(r) for r in bp.runs if r.text]
            level = 0
            if bp.paragraph_format.left_indent:
                level = int(bp.paragraph_format.left_indent / 360000)
            bullets.append(_make_bullet(bp.text, runs, level))

        if current_section.section_type in ("skills", "summary"):
            for bp in current_entry_paras:
                if bp.text.strip():
                    runs = [_run_to_textrun(r) for r in bp.runs if r.text]
                    current_section.plain_items.append(
                        _make_bullet(bp.text, runs)
                    )
        else:
            entry = ResumeEntry(
                id=str(uuid.uuid4()),
                employer=employer,
                title=title,
                start_date=start_date,
                end_date=end_date,
                location=location,
                bullets=bullets,
            )
            current_section.entries.append(entry)

        current_entry_lines = []
        current_entry_paras = []

    for para in paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if _is_heading(para, base_size):
            flush_entry()
            found_first_section = True
            heading_style: Dict[str, Any] = {
                "bold": any(r.bold for r in para.runs if r.text.strip()),
                "font_size": _para_font_size(para),
            }
            section_type = _classify_section(text)
            current_section = ResumeSection(
                id=str(uuid.uuid4()),
                heading=text,
                heading_style=heading_style,
                section_type=section_type,
            )
            sections.append(current_section)
        else:
            if not found_first_section:
                contact_lines.append(text)
            else:
                current_entry_lines.append(text)
                current_entry_paras.append(para)

    flush_entry()

    # Parse contact info from pre-section lines
    contact_info: Dict[str, str] = {}
    for line in contact_lines:
        if "@" in line:
            contact_info["email"] = line.strip()
        elif re.search(r"\d{3}[-.\s]\d{3}[-.\s]\d{4}", line):
            contact_info["phone"] = line.strip()
        elif not contact_info.get("name"):
            contact_info["name"] = line.strip()

    return ResumeIR(
        source_format="docx",
        sections=sections,
        contact_info=contact_info,
        has_tables=has_tables,
    )


# ---------------------------------------------------------------------------
# PDF Parser
# ---------------------------------------------------------------------------

def parse_pdf(file_bytes: bytes) -> ResumeIR:
    """Parse a PDF file into a ResumeIR using pdfplumber."""
    import pdfplumber

    sections: List[ResumeSection] = []
    current_section: Optional[ResumeSection] = None
    contact_lines: List[str] = []
    current_entry_lines: List[str] = []
    found_first_section = False
    all_font_sizes: List[float] = []
    contact_info: Dict[str, str] = {}

    # First pass: collect font sizes to determine base
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            chars = page.chars
            for ch in chars:
                if ch.get("size"):
                    all_font_sizes.append(ch["size"])

    if all_font_sizes:
        from collections import Counter
        counter = Counter([round(s, 1) for s in all_font_sizes])
        base_size = counter.most_common(1)[0][0]
    else:
        base_size = 11.0

    def is_heading_line(line_chars: List[Dict[str, Any]]) -> bool:
        if not line_chars:
            return False
        text = "".join(c.get("text", "") for c in line_chars).strip()
        if not text or len(text) > 70:
            return False
        if text.isupper() and len(text) < 60:
            return True
        sizes = [c.get("size", base_size) for c in line_chars]
        avg_size = sum(sizes) / len(sizes) if sizes else base_size
        bold_count = sum(
            1 for c in line_chars
            if "Bold" in (c.get("fontname", "") or "")
        )
        if bold_count > len(line_chars) * 0.5 and avg_size >= base_size:
            return True
        if avg_size > base_size * 1.1 and len(text) < 60:
            return True
        return False

    def flush_pdf_entry() -> None:
        nonlocal current_entry_lines
        if not current_section or not current_entry_lines:
            current_entry_lines = []
            return

        bullet_lines = [
            l for l in current_entry_lines
            if l.strip() and l.strip()[0] in BULLET_CHARS
        ]
        header_lines = [
            l for l in current_entry_lines
            if l.strip() and l.strip()[0] not in BULLET_CHARS
        ]

        employer, title, start_date, end_date, location = _extract_employer_title_dates(
            header_lines
        )

        if current_section.section_type in ("skills", "summary"):
            for l in current_entry_lines:
                if l.strip():
                    current_section.plain_items.append(
                        _make_bullet(l, [TextRun(text=_strip_bullet_char(l))])
                    )
        else:
            bullets = [
                _make_bullet(l, [TextRun(text=_strip_bullet_char(l))])
                for l in bullet_lines if l.strip()
            ]
            entry = ResumeEntry(
                id=str(uuid.uuid4()),
                employer=employer,
                title=title,
                start_date=start_date,
                end_date=end_date,
                location=location,
                bullets=bullets,
            )
            current_section.entries.append(entry)

        current_entry_lines = []

    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            # Group chars into lines by y-position
            lines_by_y: Dict[int, List[Dict[str, Any]]] = {}
            for ch in page.chars:
                y_key = round(ch.get("top", 0))
                lines_by_y.setdefault(y_key, []).append(ch)

            for y in sorted(lines_by_y.keys()):
                line_chars = sorted(lines_by_y[y], key=lambda c: c.get("x0", 0))
                text = "".join(c.get("text", "") for c in line_chars).strip()
                if not text:
                    continue

                if is_heading_line(line_chars):
                    flush_pdf_entry()
                    found_first_section = True
                    section_type = _classify_section(text)
                    current_section = ResumeSection(
                        id=str(uuid.uuid4()),
                        heading=text,
                        heading_style={"font_size": base_size},
                        section_type=section_type,
                    )
                    sections.append(current_section)
                else:
                    if not found_first_section:
                        contact_lines.append(text)
                        if "@" in text:
                            contact_info["email"] = text
                        elif re.search(r"\d{3}[-.\s]\d{3}[-.\s]\d{4}", text):
                            contact_info["phone"] = text
                        elif not contact_info.get("name"):
                            contact_info["name"] = text
                    else:
                        current_entry_lines.append(text)

    flush_pdf_entry()

    return ResumeIR(
        source_format="pdf",
        sections=sections,
        contact_info=contact_info,
        has_tables=False,
    )


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def parse_resume(file_bytes: bytes, filename: str) -> ResumeIR:
    """Parse a resume file (DOCX or PDF) into a ResumeIR."""
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "docx":
        return parse_docx(file_bytes)
    elif ext == "pdf":
        return parse_pdf(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only DOCX and PDF are supported.")
